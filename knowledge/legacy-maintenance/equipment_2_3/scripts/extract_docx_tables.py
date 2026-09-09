from __future__ import annotations

import csv
import json
import argparse
import hashlib
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source"
DATA = ROOT / "data"
TABLE_DIR = DATA / "tables"
OUT = ROOT / "outputs"
DOCUMENT_IDS = {}


def normalize_cell(text: str) -> str:
    return " ".join(text.split())


def find_docxs() -> list[Path]:
    candidates = sorted(SOURCE.glob("*.docx"), key=lambda p: p.name)
    if not candidates:
        raise FileNotFoundError(f"No .docx found under {SOURCE}")
    return candidates


def table_to_rows(table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([normalize_cell(cell.text) for cell in row.cells])
    return rows


def table_title(rows: list[list[str]], index: int) -> str:
    if not rows:
        return f"table_{index:02d}"
    first = [cell for cell in rows[0] if cell]
    if first:
        return " / ".join(first[:3])[:80]
    return f"table_{index:02d}"


def paragraph_outline(doc: Document) -> list[dict[str, str | int]]:
    outline: list[dict[str, str | int]] = []
    for i, para in enumerate(doc.paragraphs):
        text = normalize_cell(para.text)
        if not text:
            continue
        style = para.style.name if para.style else ""
        include = style.startswith("Heading") or style == "Caption"
        if not include:
            include = any(
                key in text
                for key in [
                    "设计示例",
                    "选型一览表",
                    "设计小结",
                    "Aspen",
                    "EDR",
                    "SW6",
                    "Column Internals",
                ]
            )
        if include:
            outline.append({"paragraph_index": i, "style": style, "text": text})
    return outline


def write_outline(doc_name: str, outline: list[dict[str, str | int]], out_json: Path, out_md: Path) -> None:
    out_json.write_text(json.dumps(outline, ensure_ascii=False, indent=2), encoding="utf-8")
    md = [
        "# DOCX 章节与证据线索",
        "",
        f"源文件：`{doc_name}`",
        "",
        "本索引只提取标题、图表题注以及含 Aspen/EDR/SW6 等证据关键词的段落，用于定位模块拆分和软件证据来源。",
        "",
        "| 段落序号 | 样式 | 文本 |",
        "| ---: | --- | --- |",
    ]
    for item in outline:
        text = str(item["text"]).replace("|", "/")
        md.append(f"| {item['paragraph_index']} | {item['style']} | {text} |")
    out_md.write_text("\n".join(md) + "\n", encoding="utf-8")


def safe_doc_id(index: int, path: Path) -> str:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    chosen = DOCUMENT_IDS.get(digest, DOCUMENT_IDS.get(path.name, "doc_" + digest[:20]))
    if not isinstance(chosen, str) or not re.fullmatch(r"[A-Za-z0-9_-]{1,100}", chosen):
        raise ValueError("Document mapping values must be safe relative IDs")
    return chosen


def main(argv=None) -> None:
    global ROOT, SOURCE, DATA, TABLE_DIR, OUT, DOCUMENT_IDS
    parser = argparse.ArgumentParser(description="Extract user-provided DOCX tables and evidence outlines")
    parser.add_argument("--source-root", required=True)
    parser.add_argument("--output-root", required=True)
    parser.add_argument("--id-mapping", help="Optional JSON filename or full file SHA256 to document ID mapping")
    args = parser.parse_args(argv)
    SOURCE = Path(args.source_root).resolve()
    ROOT = Path(args.output_root).resolve()
    DATA, TABLE_DIR, OUT = ROOT / "data", ROOT / "data/tables", ROOT / "outputs"
    DOCUMENT_IDS = json.loads(Path(args.id_mapping).read_text(encoding="utf-8")) if args.id_mapping else {}
    if not isinstance(DOCUMENT_IDS, dict): raise ValueError("ID mapping must be a JSON object")
    if ROOT == SOURCE or ROOT.is_relative_to(SOURCE) or SOURCE.is_relative_to(ROOT):
        raise ValueError("Output and source directories must be disjoint")
    if ROOT.exists() and any(ROOT.iterdir()): raise ValueError("Output must be empty")
    documents = find_docxs()
    identifiers = [safe_doc_id(i, path) for i, path in enumerate(documents)]
    if len(set(identifiers)) != len(identifiers): raise ValueError("Duplicate document ID; provide explicit unique mapping")
    ROOT.mkdir(parents=True, exist_ok=True)
    DATA.mkdir(exist_ok=True)
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(exist_ok=True)

    all_tables = []
    all_outlines = []
    table_md = [
        "# DOCX 表格提取总索引",
        "",
        "| 文档ID | 源文件 | 表序号 | 标题线索 | 行数 | 列数 | CSV |",
        "| --- | --- | ---: | --- | ---: | ---: | --- |",
    ]
    outline_md = [
        "# DOCX 章节与证据总索引",
        "",
        "| 文档ID | 源文件 | 段落序号 | 样式 | 文本 |",
        "| --- | --- | ---: | --- | --- |",
    ]

    for doc_index, docx_path in enumerate(find_docxs()):
        doc_id = safe_doc_id(doc_index, docx_path)
        doc_table_dir = TABLE_DIR / doc_id
        doc_table_dir.mkdir(parents=True, exist_ok=True)
        doc = Document(str(docx_path))
        doc_index_rows = []

        for i, table in enumerate(doc.tables):
            rows = table_to_rows(table)
            csv_path = doc_table_dir / f"table_{i:02d}.csv"
            with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerows(rows)
            row = {
                "doc_id": doc_id,
                "source_file": docx_path.name,
                "index": i,
                "title_hint": table_title(rows, i),
                "rows": len(rows),
                "columns": max((len(row) for row in rows), default=0),
                "csv": str(csv_path.relative_to(ROOT)),
            }
            doc_index_rows.append(row)
            all_tables.append(row)
            table_md.append(
                f"| {doc_id} | {docx_path.name} | {i} | {row['title_hint']} | {row['rows']} | {row['columns']} | `{row['csv']}` |"
            )

        doc_outline = paragraph_outline(doc)
        for item in doc_outline:
            full_item = {"doc_id": doc_id, "source_file": docx_path.name, **item}
            all_outlines.append(full_item)
            text = str(item["text"]).replace("|", "/")
            outline_md.append(f"| {doc_id} | {docx_path.name} | {item['paragraph_index']} | {item['style']} | {text} |")

        (DATA / f"tables_index_{doc_id}.json").write_text(
            json.dumps(doc_index_rows, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        write_outline(
            docx_path.name,
            doc_outline,
            DATA / f"docx_outline_{doc_id}.json",
            OUT / f"DOCX章节与证据线索_{doc_id}.md",
        )

    (DATA / "tables_index.json").write_text(json.dumps(all_tables, ensure_ascii=False, indent=2), encoding="utf-8")
    (DATA / "docx_outline.json").write_text(json.dumps(all_outlines, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "表格提取索引.md").write_text("\n".join(table_md) + "\n", encoding="utf-8")
    (OUT / "DOCX章节与证据线索.md").write_text("\n".join(outline_md) + "\n", encoding="utf-8")

    print(f"extracted {len(all_tables)} tables from {len(find_docxs())} docx files")
    print(f"extracted {len(all_outlines)} outline/evidence paragraphs")
    print(DATA / "tables_index.json")
    print(DATA / "docx_outline.json")
    print(OUT / "表格提取索引.md")
    print(OUT / "DOCX章节与证据线索.md")


if __name__ == "__main__":
    main()
